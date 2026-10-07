"""Reconciles the 9 yearly SF3 spreadsheets in data/raw/ into one tidy
table: (ecode, local_authority, year, measure, value). Sheet layout and
header position both vary by year, so headers are found by searching for
the "Local Authority" cell rather than assumed by row/column number."""

import re
from pathlib import Path

import duckdb
import openpyxl
import pandas as pd

RAW_DIR = Path("data/raw")
OUT_PATH = Path("data/processed/lgps_sf3_tidy.parquet")

FILES_BY_YEAR = {
    "2016-17": "LA_drop_down.xlsx",
    "2017-18": "LA_drop_down_revised.xlsx",
    "2018-19": "LA_drop_down_2018-19.xlsx",
    "2019-20": "LA_drop_down_2019-20_revised.xlsx",
    "2020-21": "LA_drop_down_2020-21_revised.xlsx",
    "2021-22": "LA_drop_down_2021-22_April_update.xlsx",
    "2022-23": "LA_drop_down_2022-23_-_ecomms_-_July_2024.xlsx",
    "2023-24": "LA_drop_down_2023-24_-_June_2025_-_ecomms.xlsx",
    "2024-25": "LA_drop_down_2024-25_-_ecomms.xlsx",
}

EXP_INC_MEASURES = [
    "pension", "lump_retire", "lump_opt", "lump_death", "othben",
    "transf_out", "penprem", "mgmtexp", "othexp", "totpens_exp",
    "contrib_empee", "contrib_emper", "invinc", "transf_in", "othinc",
    "totpens_inc",
]

ECODE_RE = re.compile(r"^(EW\d{3}|E\d{4}|W\d{4}|G\d{3}(-\d+)?|E9999|W9999)$")


def find_header_row(ws, search_rows=15, search_cols=10):
    """Find the row containing the "Local Authority" column header."""
    for r in range(1, search_rows + 1):
        for c in range(1, search_cols + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip().lower() == "local authority":
                return r
    raise ValueError(f"Could not find 'Local Authority' header in {ws.title}")


def cell_text(ws, row, col):
    v = ws.cell(row, col).value
    return v.strip() if isinstance(v, str) and v.strip() else None


def get_label(ws, header_row, col):
    """Label on the header row itself, or two rows above for old-format sheets."""
    return cell_text(ws, header_row, col) or cell_text(ws, header_row - 2, col)


ID_LABELS = {
    "no.", "local authority", "ecodes", "ecode", "e code", "ons code",
    "notes", "mhclg e code", "delta code", "delta codes",
}


def find_id_columns(ws, header_row, max_col):
    """Locate the fund-name and fund-code columns on the header row."""
    name_col = ecode_col = None
    for c in range(1, max_col + 1):
        v = cell_text(ws, header_row, c)
        if not v:
            continue
        if v.lower() == "local authority":
            name_col = c
        elif "e code" in v.lower() or v.lower() in ("ecodes", "ecode"):
            ecode_col = c
    if not name_col or not ecode_col:
        raise ValueError(f"Could not locate name/ecode columns in {ws.title}")
    return name_col, ecode_col


def find_measure_columns(ws, header_row, max_col):
    """All non-identifier columns with a label, left to right."""
    cols = []
    for c in range(1, max_col + 1):
        label = get_label(ws, header_row, c)
        if label and label.lower() not in ID_LABELS:
            cols.append(c)
    return cols


def iter_fund_rows(ws, header_row, ecode_col, max_col):
    """Yield (row, ecode) for each real fund or national-total row."""
    for r in range(header_row + 1, ws.max_row + 1):
        ecode = cell_text(ws, r, ecode_col)
        if ecode and ECODE_RE.match(ecode):
            yield r, ecode


def extract_exp_inc(wb, year):
    """Costs, contributions and investment income for one year, tidy."""
    sheet = "Data_Exp_and_Inc" if "Data_Exp_and_Inc" in wb.sheetnames else "Data"
    ws = wb[sheet]
    header_row = find_header_row(ws)
    name_col, ecode_col = find_id_columns(ws, header_row, max_col=10)
    measure_cols = find_measure_columns(ws, header_row, max_col=30)

    if len(measure_cols) != len(EXP_INC_MEASURES):
        raise ValueError(
            f"{year}: expected {len(EXP_INC_MEASURES)} measure columns on "
            f"'{sheet}', found {len(measure_cols)}"
        )

    records = []
    for r, ecode in iter_fund_rows(ws, header_row, ecode_col, max_col=max(measure_cols)):
        fund_name = cell_text(ws, r, name_col)
        for measure, col in zip(EXP_INC_MEASURES, measure_cols):
            value = ws.cell(r, col).value
            if value is None:
                continue
            records.append((ecode, fund_name, year, measure, value))

    return pd.DataFrame(
        records, columns=["ecode", "local_authority", "year", "measure", "value"]
    )


MEMBERSHIP_PATTERNS = {
    "empler_tot": ("number of employers", "total number of employers"),
    "contmem_tot": ("number of contributing members", "total number of contributing members"),
    "pensioner_tot": ("number of pensioners", "total number of pensioners"),
    "defmemb_tot": ("former members", "total number of former members"),
    "totmember_tot": ("total number of members", "total number of members"),
}


def find_label_columns(ws, header_row, contains, max_col):
    contains = contains.lower()
    cols = []
    for c in range(1, max_col + 1):
        label = get_label(ws, header_row, c)
        if label and contains in label.lower():
            cols.append(c)
    return cols


def find_exact_label_column(ws, header_row, text, max_col):
    text = text.lower()
    return [
        c for c in range(1, max_col + 1)
        if (get_label(ws, header_row, c) or "").lower() == text
    ]


def resolve_total_column(ws, header_row, broad, total_phrase, max_col):
    """total_phrase must match exactly; broad plus the 'Total' sub-header is the fallback."""
    total_cols = find_exact_label_column(ws, header_row, total_phrase, max_col)
    if len(total_cols) == 1:
        return total_cols[0]

    cols = find_label_columns(ws, header_row, broad, max_col)
    totals = [c for c in cols if (cell_text(ws, header_row + 1, c) or "").lower() == "total"]
    if len(totals) == 1:
        return totals[0]
    raise ValueError(
        f"Could not resolve a unique '{total_phrase}' column in {ws.title} "
        f"(broad match: {len(cols)}, total sub-header match: {len(totals)})"
    )


def extract_by_label(ws, year, id_cols, measure_patterns):
    header_row = find_header_row(ws)
    name_col, ecode_col = id_cols(ws, header_row)
    max_col = ws.max_column

    resolved = {
        code: resolve_total_column(ws, header_row, broad, total_phrase, max_col)
        for code, (broad, total_phrase) in measure_patterns.items()
    }

    records = []
    for r, ecode in iter_fund_rows(ws, header_row, ecode_col, max_col):
        fund_name = cell_text(ws, r, name_col)
        for code, col in resolved.items():
            value = ws.cell(r, col).value
            if value is None:
                continue
            records.append((ecode, fund_name, year, code, value))

    return pd.DataFrame(
        records, columns=["ecode", "local_authority", "year", "measure", "value"]
    )


def extract_market_value(ws, year, id_cols):
    header_row = find_header_row(ws)
    name_col, ecode_col = id_cols(ws, header_row)
    max_col = ws.max_column

    mkt_cols = find_label_columns(ws, header_row, "market value of the fund", max_col)
    if len(mkt_cols) != 2:
        raise ValueError(
            f"{year}: expected 2 'market value of the fund' columns in "
            f"{ws.title}, found {len(mkt_cols)}"
        )
    coded = {}
    for c in mkt_cols:
        label = get_label(ws, header_row, c).lower()
        coded["mktval_startyr" if "1 april" in label else "mktval_endyr"] = c

    records = []
    for r, ecode in iter_fund_rows(ws, header_row, ecode_col, max_col):
        fund_name = cell_text(ws, r, name_col)
        for code, col in coded.items():
            value = ws.cell(r, col).value
            if value is None:
                continue
            records.append((ecode, fund_name, year, code, value))

    return pd.DataFrame(
        records, columns=["ecode", "local_authority", "year", "measure", "value"]
    )


def default_id_cols(ws, header_row):
    return find_id_columns(ws, header_row, max_col=10)


def extract_membership_and_assets(wb, year):
    """Membership totals and fund market value for one year, tidy."""
    if "Data2" in wb.sheetnames:
        ws = wb["Data2"]
        membership = extract_by_label(ws, year, default_id_cols, MEMBERSHIP_PATTERNS)
        assets = extract_market_value(ws, year, default_id_cols)
    else:
        membership = extract_by_label(
            wb["Data_Memo_SectionA"], year, default_id_cols, MEMBERSHIP_PATTERNS
        )
        assets = extract_market_value(
            wb["Data_Memo_SectionB_to_F"], year, default_id_cols
        )
    return pd.concat([membership, assets], ignore_index=True)


def main():
    frames = []
    for year, fname in FILES_BY_YEAR.items():
        path = RAW_DIR / fname
        print(f"{year}: reading {fname}")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        frames.append(extract_exp_inc(wb, year))
        frames.append(extract_membership_and_assets(wb, year))

    tidy = pd.concat(frames, ignore_index=True)
    tidy["value"] = pd.to_numeric(tidy["value"], errors="coerce")
    tidy = tidy.dropna(subset=["value"])

    print(f"\n{len(tidy):,} rows, {tidy['measure'].nunique()} measures, "
          f"{tidy['year'].nunique()} years, {tidy['ecode'].nunique()} fund codes")

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    duckdb.sql(f"COPY tidy TO '{OUT_PATH}' (FORMAT PARQUET)")
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
